"""读文本文件；pdf/docx/xlsx 抽成文本。短的内联，长的落盘再 grep。"""

from __future__ import annotations

import hashlib
import os
from pathlib import Path
from typing import Any, Callable, Dict, Optional

NAME = "read"
REPLAY = "safe"
READ_ONLY = True

MAX_LINES = 1000
MAX_BYTES = 100_000
# 长文档抽完超过这个长度就落盘，让模型 grep，而不是按行翻二进制原文。
SPILL_CHARS = 8000
PREVIEW_LINES = 30
PREVIEW_CHARS = 2000


def resolve_path(path: str, workspace: str) -> Path:
    """相对路径拼到 workspace；绝对路径原样。不限制出 workspace。"""
    root = Path(workspace or os.getcwd()).resolve()
    raw = Path(path)
    return (raw if raw.is_absolute() else root / raw).resolve()


def before_tool_deny_empty_path(event: Dict[str, Any]):
    path = ((event.get("args") or {}).get("path") or "").strip()
    if not path:
        return {"block": {"reason": "path 为空"}}
    return None


def after_tool_truncate_output(event: Dict[str, Any]):
    content = event.get("content") or ""
    limit = 8000
    if len(content) <= limit:
        return None
    return {"content": content[:limit] + "\n...[output truncated]"}


BEFORE_HOOKS = [before_tool_deny_empty_path]
# read 只按 offset 翻页，不再 after 切 8k（那会切掉续读提示）。
# 函数留给 grep 等工具继续当输出封顶。
AFTER_HOOKS: list = []


async def execute(args: Dict[str, Any], sandbox=None, workspace=None, session_dir=None) -> str:
    path = (args.get("path") or "").strip()
    if not path:
        raise ValueError("path 为空")
    offset = args.get("offset")
    limit = args.get("limit")
    root = workspace or os.getcwd()
    file_path = resolve_path(path, root)
    if not file_path.exists():
        raise FileNotFoundError(f"file not found: {path}")
    if file_path.is_dir():
        raise IsADirectoryError(f"is a directory: {path}")
    extractor = EXTRACTORS.get(file_path.suffix.lower())
    if extractor is not None:
        return _read_extracted(
            file_path,
            extractor,
            workspace=root,
            session_dir=session_dir,
            offset=offset,
            limit=limit,
        )
    return _read_text(file_path, offset=offset, limit=limit)


def _read_extracted(
    file_path: Path,
    extractor: Callable[[Path], str],
    *,
    workspace: str,
    session_dir: Optional[str],
    offset=None,
    limit=None,
) -> str:
    body = extractor(file_path).strip()
    if not body:
        raise ValueError(
            f"no extractable text in {file_path.name}; it may be scanned or empty"
        )
    kind = file_path.suffix.lower().lstrip(".")
    full = f"[{kind}] {file_path.name}\n\n{body}"
    if session_dir and len(full) > SPILL_CHARS:
        saved = _write_extract(session_dir, workspace, file_path, full)
        return _stub(file_path.name, saved, full)
    return _paginate(full, offset=offset, limit=limit)


def _write_extract(session_dir: str, workspace: str, file_path: Path, full: str) -> str:
    key = hashlib.sha256(str(file_path.resolve()).encode("utf-8")).hexdigest()[:12]
    dest = Path(session_dir) / "workspace" / "tools_result" / "read" / f"{key}.txt"
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(full, encoding="utf-8")
    root = Path(workspace or ".").resolve()
    try:
        return dest.resolve().relative_to(root).as_posix()
    except ValueError:
        return str(dest)


def _stub(name: str, saved: str, full: str) -> str:
    lines = full.splitlines()
    preview = "\n".join(lines[:PREVIEW_LINES])
    if len(preview) > PREVIEW_CHARS:
        preview = preview[:PREVIEW_CHARS].rstrip() + "\n..."
    elif len(lines) > PREVIEW_LINES:
        preview = preview.rstrip() + "\n..."
    return (
        f"[read] {name}\n"
        f"saved: {saved}\n"
        f"chars: {len(full)}\n\n"
        f"{preview}\n\n"
        "To find something, grep or read the saved path "
        "(pass path= that file). Do not parse the original with bash."
    )


def _extract_pdf(file_path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(file_path))
    parts = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if not text:
            continue
        parts.append(f"--- page {i} ---\n{text}")
    return "\n\n".join(parts)


def _extract_docx(file_path: Path) -> str:
    from docx import Document

    doc = Document(str(file_path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_xlsx(file_path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(file_path), read_only=True, data_only=True)
    try:
        chunks = []
        for name in wb.sheetnames:
            chunks.append(f"--- sheet {name} ---")
            ws = wb[name]
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c) for c in row]
                if any(cells):
                    chunks.append("\t".join(cells))
        return "\n".join(chunks)
    finally:
        wb.close()


EXTRACTORS: Dict[str, Callable[[Path], str]] = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
}


def _paginate(text: str, offset=None, limit=None) -> str:
    lines = text.splitlines()
    total = len(lines)
    start = 0
    if offset is not None:
        start = max(int(offset) - 1, 0)
    if start >= total:
        raise ValueError(f"offset {offset} is beyond end of file ({total} lines)")
    end = total
    if limit is not None:
        end = min(start + max(int(limit), 0), total)
    selected = lines[start:end]
    encoded = "\n".join(selected).encode("utf-8")
    if len(encoded) > MAX_BYTES:
        cut = selected
        while cut and len("\n".join(cut).encode("utf-8")) > MAX_BYTES:
            cut = cut[:-1]
        selected = cut
        end = start + len(selected)
    if len(selected) > MAX_LINES:
        selected = selected[:MAX_LINES]
        end = start + len(selected)
    numbered = []
    for i, line in enumerate(selected, start=start + 1):
        numbered.append(f"{i:6d}|{line}")
    body = "\n".join(numbered)
    shown_start = start + 1
    shown_end = start + len(selected)
    footer = f"\n\n[Showing lines {shown_start}-{shown_end} of {total}]"
    if shown_end < total:
        footer += f" Use offset={shown_end + 1} to continue."
    return body + footer


def _read_text(file_path: Path, offset=None, limit=None) -> str:
    data = file_path.read_bytes()
    if b"\x00" in data[:8000]:
        raise ValueError(f"binary file cannot be read as text: {file_path.name}")
    text = data.decode("utf-8", errors="replace")
    return _paginate(text, offset=offset, limit=limit)
