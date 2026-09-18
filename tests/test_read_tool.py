"""read 工具：行号、offset/limit、路径解析。"""
from __future__ import annotations

import asyncio
import json
from pathlib import Path

import pytest

from agent_loop.runtime.tool_runtime import ToolRuntime
from agent_loop.tools.read_tool import execute, resolve_path


def _write(tmp_path: Path, name: str, text: str) -> Path:
    p = tmp_path / name
    p.write_text(text, encoding="utf-8")
    return p


def test_resolve_path_allows_outside(tmp_path):
    got = resolve_path("../secret", str(tmp_path))
    assert got == (tmp_path / ".." / "secret").resolve()


def test_read_outside_workspace(tmp_path):
    workspace = tmp_path / "ws"
    other = tmp_path / "other"
    workspace.mkdir()
    other.mkdir()
    (other / "ref.py").write_text("x = 1\n", encoding="utf-8")
    out = asyncio.run(
        execute({"path": str(other / "ref.py")}, workspace=str(workspace))
    )
    assert "     1|x = 1" in out


def test_read_numbers_lines(tmp_path):
    _write(tmp_path, "a.py", "import os\ndef main():\n    pass\n")
    out = asyncio.run(execute({"path": "a.py"}, workspace=str(tmp_path)))
    assert "     1|import os" in out
    assert "     2|def main():" in out
    assert "Showing lines 1-3 of 3" in out


def test_read_offset_limit(tmp_path):
    _write(tmp_path, "a.txt", "a\nb\nc\nd\n")
    out = asyncio.run(
        execute({"path": "a.txt", "offset": 2, "limit": 2}, workspace=str(tmp_path))
    )
    assert "     2|b" in out
    assert "     3|c" in out
    assert "     1|a" not in out
    assert "Use offset=4 to continue" in out


def test_read_missing_file(tmp_path):
    with pytest.raises(FileNotFoundError):
        asyncio.run(execute({"path": "nope.txt"}, workspace=str(tmp_path)))


def test_read_offset_beyond_end(tmp_path):
    _write(tmp_path, "a.txt", "only\n")
    with pytest.raises(ValueError, match="beyond end"):
        asyncio.run(execute({"path": "a.txt", "offset": 9}, workspace=str(tmp_path)))


def test_short_pdf_returns_inline(tmp_path, monkeypatch):
    from agent_loop.tools import read_tool

    monkeypatch.setitem(read_tool.EXTRACTORS, ".pdf", lambda _p: "Invoice UZR2GULN $100")
    (tmp_path / "inv.pdf").write_bytes(b"%PDF-\x00fake")
    out = asyncio.run(execute({"path": "inv.pdf"}, workspace=str(tmp_path)))
    assert "Invoice UZR2GULN $100" in out
    assert "     1|[pdf] inv.pdf" in out
    assert "saved:" not in out


def test_long_pdf_spills_to_session(tmp_path, monkeypatch):
    from agent_loop.tools import read_tool

    body = ("lorem ipsum dolor sit amet\n" * 400)
    assert len(f"[pdf] big.pdf\n\n{body}") > read_tool.SPILL_CHARS
    monkeypatch.setitem(read_tool.EXTRACTORS, ".pdf", lambda _p: body)
    (tmp_path / "big.pdf").write_bytes(b"%PDF-\x00fake")
    session = tmp_path / "chat"
    out = asyncio.run(
        execute(
            {"path": "big.pdf"},
            workspace=str(tmp_path),
            session_dir=str(session),
        )
    )
    assert "[read] big.pdf" in out
    assert "saved:" in out
    assert "chars:" in out
    assert "grep or read the saved path" in out
    saved_line = [ln for ln in out.splitlines() if ln.startswith("saved:")][0]
    saved = saved_line.split(" ", 1)[1]
    text = Path(saved).read_text(encoding="utf-8") if Path(saved).is_absolute() else (
        tmp_path / saved
    ).read_text(encoding="utf-8")
    assert "lorem ipsum" in text
    assert body.strip() in text


def test_long_pdf_without_session_paginates(tmp_path, monkeypatch):
    from agent_loop.tools import read_tool

    body = "row\n" * 5000
    monkeypatch.setitem(read_tool.EXTRACTORS, ".pdf", lambda _p: body)
    (tmp_path / "big.pdf").write_bytes(b"%PDF-fake")
    out = asyncio.run(execute({"path": "big.pdf"}, workspace=str(tmp_path)))
    assert "saved:" not in out
    assert "Use offset=" in out


def test_empty_pdf_errors(tmp_path):
    from pypdf import PdfWriter

    blank = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(blank)
    with pytest.raises(ValueError, match="no extractable text"):
        asyncio.run(execute({"path": "blank.pdf"}, workspace=str(tmp_path)))


def test_docx_extracts_paragraphs(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("Hello docx")
    path = tmp_path / "note.docx"
    doc.save(path)
    out = asyncio.run(execute({"path": "note.docx"}, workspace=str(tmp_path)))
    assert "Hello docx" in out
    assert "saved:" not in out


def test_xlsx_extracts_cells(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["A1"] = "sku"
    ws["B1"] = "qty"
    ws["A2"] = "ABC"
    ws["B2"] = 3
    path = tmp_path / "book.xlsx"
    wb.save(path)
    out = asyncio.run(execute({"path": "book.xlsx"}, workspace=str(tmp_path)))
    assert "--- sheet Sales ---" in out
    assert "ABC" in out
    assert "saved:" not in out


def test_read_pages_by_offset_not_8k_hook(tmp_path):
    lines = [f"line-{i:04d} " + ("x" * 40) for i in range(1, 1501)]
    _write(tmp_path, "big.txt", "\n".join(lines) + "\n")
    runtime = ToolRuntime()
    runtime.workspace = str(tmp_path)
    first = asyncio.run(
        runtime.call_tool(
            {
                "id": "c1",
                "function": {"name": "read", "arguments": '{"path": "big.txt"}'},
            }
        )
    )
    assert first.is_error is False
    assert "...[output truncated]" not in first.content
    assert "Use offset=1001 to continue" in first.content
    assert "line-1000" in first.content
    assert "line-1001" not in first.content
    assert len(first.content) > 8000
    second = asyncio.run(
        runtime.call_tool(
            {
                "id": "c2",
                "function": {
                    "name": "read",
                    "arguments": '{"path": "big.txt", "offset": 1001}',
                },
            }
        )
    )
    assert "line-1001" in second.content
    assert "Showing lines 1001-1500 of 1500" in second.content


def test_runtime_read_ok(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "hello.txt").write_text("hi\n", encoding="utf-8")
    runtime = ToolRuntime()
    runtime.workspace = str(tmp_path)
    result = asyncio.run(
        runtime.call_tool(
            {
                "id": "c1",
                "function": {
                    "name": "read",
                    "arguments": '{"path": "hello.txt"}',
                },
            }
        )
    )
    assert result.is_error is False
    assert "     1|hi" in result.content


def _min_png() -> bytes:
    import struct
    import zlib

    def chunk(tag: bytes, data: bytes) -> bytes:
        crc = zlib.crc32(tag + data) & 0xFFFFFFFF
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", crc)

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    raw = zlib.compress(b"\x00\xff\x00\x00")
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", raw) + chunk(b"IEND", b"")


def test_sniff_image_png_jpeg_webp():
    from agent_loop.tools.read_tool import sniff_image

    assert sniff_image(_min_png())[0] == "image/png"
    assert sniff_image(b"\xff\xd8\xff\xe0xxxx")[0] == "image/jpeg"
    assert sniff_image(b"RIFF\x00\x00\x00\x00WEBP....")[0] == "image/webp"
    assert sniff_image(b"not an image") is None


def test_read_png_returns_media_and_copies(tmp_path):
    from agent_loop.runtime.records import ToolOutput

    png = tmp_path / "shot.png"
    png.write_bytes(_min_png())
    session = tmp_path / "chat"
    out = asyncio.run(
        execute(
            {"path": "shot.png"},
            workspace=str(tmp_path),
            session_dir=str(session),
        )
    )
    assert isinstance(out, ToolOutput)
    assert "Read image file [image/png]" in out.content
    assert out.media["kind"] == "image"
    assert out.media["mime"] == "image/png"
    saved = Path(out.media["path"])
    assert saved.is_file()
    assert saved.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n"


def test_read_image_jsonl_has_media_not_base64(tmp_path):
    from agent_loop.session_log import SessionLog

    png = tmp_path / "a.png"
    png.write_bytes(_min_png())
    log_path = tmp_path / "session.jsonl"
    runtime = ToolRuntime()
    runtime.workspace = str(tmp_path)
    runtime.session_dir = str(tmp_path / "sess")
    runtime.attach_log(SessionLog(log_path))
    result = asyncio.run(
        runtime.call_tool(
            {
                "id": "c1",
                "function": {"name": "read", "arguments": '{"path": "a.png"}'},
            }
        )
    )
    assert result.media["kind"] == "image"
    assert "base64" not in result.content
    rows = SessionLog(log_path).read_all()
    attached = [r for r in rows if r.get("type") == "image_attached"]
    assert len(attached) == 1
    assert attached[0]["kind"] == "record"
    assert attached[0]["name"].endswith(".png")
    assert "base64" not in json.dumps(attached[0])


def test_entries_to_messages_attaches_image_on_user_not_tool(tmp_path):
    from agent_loop.recover import entries_to_messages
    from agent_loop.tools.read_tool import IMAGE_OMIT_NOTE, encode_image_data_url

    png = tmp_path / "a.png"
    png.write_bytes(_min_png())
    row = {
        "kind": "entry",
        "type": "tool_result",
        "seq": 1,
        "tool_call_id": "call1",
        "tool_name": "read",
        "content": "Read image file [image/png]\nsaved: x",
        "media": {"kind": "image", "mime": "image/png", "path": str(png)},
    }
    with_img = entries_to_messages([row], supports_images=True)
    assert with_img[0]["role"] == "tool"
    assert with_img[0]["content"].startswith("Read image file")
    assert with_img[1]["role"] == "user"
    blocks = with_img[1]["content"]
    assert blocks[0]["type"] == "text"
    assert blocks[1]["type"] == "image_url"
    url = blocks[1]["image_url"]["url"]
    assert url.startswith("data:image/png;base64,")
    assert encode_image_data_url(str(png), "image/png") == url

    no_img = entries_to_messages([row], supports_images=False)
    assert len(no_img) == 1
    assert no_img[0]["role"] == "tool"
    assert IMAGE_OMIT_NOTE in no_img[0]["content"]
    assert "base64" not in no_img[0]["content"]
